from pathlib import Path
import json,subprocess,tempfile,unittest
from datetime import datetime
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1]
class Finance(unittest.TestCase):
 def test_recovered_financial_formulas(self):
  spec=json.loads((ROOT/'data/plantilla.json').read_text());wb=Workbook();wb.remove(wb.active)
  for definition in spec['sheets']:
   ws=wb.create_sheet(definition['name'])
   for c in definition['cells']:
    value=c['v'];value=datetime.fromisoformat(value['date']) if isinstance(value,dict) and 'date' in value else value
    ws.cell(c['r'],c['c'],value)
  wb['Ajustes']['B6']=datetime(2026,10,1);wb['Ajustes']['B12']=datetime(2026,10,1);wb['Cuentas']['B9']=datetime(2026,10,27)
  for r,kind,category,amount in [(6,'Ingreso','Ingresos',350),(7,'Aporte fondo','Fondo gasolina',10),(8,'Gasto','Fondo gasolina',4),(9,'Aporte ahorro','Ahorro general',175)]:
   for c,v in enumerate([datetime(2026,10,26),kind,category,'Solo datos de prueba','USD',amount,''],1):wb['Registro'].cell(r,c,v)
   for col,formula in spec['registerTemplates'].items():wb['Registro'][f'{col}{r}']=formula.replace('{r}',str(r))
  with tempfile.TemporaryDirectory(prefix='hogar-finance-recovery-') as tmp:
   p=Path(tmp);(p/'source').mkdir();(p/'result').mkdir();wb.save(p/'source/test.xlsx');wb.close()
   subprocess.run(['soffice','-env:UserInstallation=file://'+str(p/'profile'),'--headless','--convert-to','xlsx','--outdir',str(p/'result'),str(p/'source/test.xlsx')],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=90)
   out=load_workbook(p/'result/test.xlsx',data_only=True)
   self.assertEqual(out['Registro']['K8'].value,'OK');self.assertEqual(out['Fondos']['F11'].value,6);self.assertEqual(out['Cuentas']['F6'].value,346);self.assertEqual(out['Inicio']['B15'].value,0)
   self.assertEqual(out['Mensual']['D6'].value,10);self.assertEqual(out['Mensual']['I6'].value,4)
   guava=next(c['r'] for d in spec['sheets'] if d['name']=='Lista compras' for c in d['cells'] if c['c']==2 and c['v']=='Guayaba');self.assertAlmostEqual(out['Lista compras'][f'H{guava}'].value,3.864)
   out.close()
if __name__=='__main__':unittest.main()
